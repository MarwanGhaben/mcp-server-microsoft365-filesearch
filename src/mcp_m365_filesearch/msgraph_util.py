import os
import requests
import time
import json
import mimetypes
from io import BytesIO
from docx import Document
import openpyxl
from llama_index.core import SimpleDirectoryReader
from .logger_config import setup_logger

# Initialize logger
logger = setup_logger()

GRAPH_URL = "https://graph.microsoft.com/v1.0"

SITE_NAME_TO_ID = {
    "Mazoo": "marwanmostafa.sharepoint.com,121f66a5-f7c5-4f4b-839a-74bd313275e4,78f6f561-fcb0-4138-bef1-7f119aabc8aa"
}

# ----------------------
# Graph Search
# ----------------------
def search_graph(query_text, access_token, region, size=20, from_index=0):
    logger.info(f"Searching Microsoft Graph for query: {query_text} (from={from_index}, size={size})")
    headers = {
        "Authorization": f"Bearer {access_token}",
        "Content-Type": "application/json"
    }
    url = f"{GRAPH_URL}/search/query"
    body = {
        "requests": [
            {
                "entityTypes": ["driveItem"],
                "query": {"queryString": query_text},
                "fields": [
                    "name", "webUrl", "id", "parentReference",
                    "createdBy", "createdDateTime", "lastModifiedBy", "lastModifiedDateTime"
                ],
                "from": from_index,
                "size": size,
                "region": region
            }
        ]
    }
    response = requests.post(url, headers=headers, json=body)
    if response.status_code == 200:
        logger.info("Search completed successfully.")
        return response.json()
    else:
        logger.error(f"Search failed: {response.status_code} - {response.text}")
        return None

# ----------------------
# Response Parser
# ----------------------
def parse_search_response(search_results, file_type, file_extension):
    results = []
    hits_containers = search_results.get("value", [])
    if hits_containers and isinstance(hits_containers, list):
        hits = hits_containers[0].get("hitsContainers", [])
        if hits and isinstance(hits, list):
            for result in hits[0].get("hits", []):
                resource = result.get("resource", {})
                az_search_rank = result.get("rank")
                summary = result.get("summary", "")
                file_name = resource.get("name", "")
                file_url = resource.get("webUrl")

                if file_name and (
                    file_type == "all" or any(file_name.endswith(f".{ext}") for ext in file_extension)
                ):
                    results.append({
                        "name": file_name,
                        "url": file_url,
                        "summary": summary,
                        "rank": az_search_rank,
                        "source": classify_source(file_url),
                        "created_by": resource.get("createdBy", {}).get("user", {}),
                        "created_date": resource.get("createdDateTime"),
                        "last_modified_by": resource.get("lastModifiedBy", {}).get("user", {}),
                        "last_modified_date": resource.get("lastModifiedDateTime"),
                        "fileid": resource.get("id"),
                        "parent_reference": resource.get("parentReference", {}),
                        "drive_id": resource.get("parentReference", {}).get("driveId"),
                    })
    return results

# ----------------------
# Download Helpers
# ----------------------
def classify_source(web_url):
    logger.debug(f"Classifying source for URL: {web_url}")
    if "my.sharepoint.com/personal/" in web_url:
        return "OneDrive"
    return "SharePoint"

async def download_file(drive_id, item_id, access_token, offset=0, limit=50):
    logger.info(f"Downloading file with ID: {item_id} from drive: {drive_id}")
    headers = {"Authorization": f"Bearer {access_token}"}
    metadata_url = f"{GRAPH_URL}/drives/{drive_id}/items/{item_id}"

    metadata_response = requests.get(metadata_url, headers=headers)
    if metadata_response.status_code == 200:
        metadata = metadata_response.json()
        file_name = metadata.get("name", f"{item_id}.bin")
    else:
        logger.error(f"Failed to fetch metadata: {metadata_response.status_code} - {metadata_response.text}")
        return None

    current_dir = os.path.dirname(os.path.abspath(__file__))
    local_dir = os.path.join(current_dir, ".local")
    item_folder = os.path.join(local_dir, "downloads", drive_id, item_id)
    os.makedirs(item_folder, exist_ok=True)

    existing_files = os.listdir(item_folder)
    if existing_files:
        existing_file_path = os.path.join(item_folder, existing_files[0])
        file_age = time.time() - os.path.getmtime(existing_file_path)
        if file_age < 24 * 3600:
            logger.info(f"Using cached file: {existing_file_path}")
            return await _read_file_content(existing_file_path, offset=offset, limit=limit)
        else:
            logger.info(f"Deleting old file: {existing_file_path}")
            os.remove(existing_file_path)

    file_path = os.path.join(item_folder, file_name)
    content_url = f"{metadata_url}/content"

    response = requests.get(content_url, headers=headers, stream=True)
    if response.status_code == 200:
        with open(file_path, "wb") as f:
            for chunk in response.iter_content(chunk_size=8192):
                f.write(chunk)
        logger.info(f"File downloaded: {file_path}")
        return await _read_file_content(file_path, offset=offset, limit=limit)
    else:
        logger.error(f"Download failed: {response.status_code} - {response.text}")
        return None

# ----------------------
# File Reader + Pagination
# ----------------------
async def _read_file_content(file_path, offset=0, limit=50):
    try:
        cache_path = f"{file_path}.cache.json"
        if os.path.exists(cache_path):
            age = time.time() - os.path.getmtime(cache_path)
            if age < 24 * 3600:
                logger.info(f"Using cached content: {cache_path}")
                with open(cache_path, "r", encoding="utf-8") as f:
                    return json.load(f)

        try:
            reader = SimpleDirectoryReader(input_files=[file_path])
            docs = reader.load_data()
            if docs:
                serialized = [{"text": doc.text, **doc.metadata} for doc in docs]
                with open(cache_path, "w", encoding="utf-8") as f:
                    json.dump(serialized, f, ensure_ascii=False, indent=4)
                return serialized
        except Exception as e:
            logger.warning(f"LlamaIndex failed, trying fallback: {e}")

        text_output = ""
        if file_path.endswith(".docx"):
            doc = Document(file_path)
            text_output = "\n".join([p.text for p in doc.paragraphs])

        elif file_path.endswith(".xlsx"):
            wb = openpyxl.load_workbook(file_path, data_only=True)
            chunks = []
            for sheet in wb.worksheets:
                chunks.append(f"--- Sheet: {sheet.title} ---")
                start = offset
                end = offset + limit
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i < start:
                        continue
                    if i >= end:
                        chunks.append(f"[... {limit} rows shown. Use offset={offset + limit} to continue ...]")
                        break
                    row_text = [str(cell) if cell else "" for cell in row]
                    chunks.append(" | ".join(row_text))
            text_output = "\n".join(chunks)

        if text_output:
            result = [{"text": text_output, "source": "manual"}]
            with open(cache_path, "w", encoding="utf-8") as f:
                json.dump(result, f, ensure_ascii=False, indent=4)
            return result
        else:
            return [{"text": "[No readable text found]", "source": "manual"}]

    except Exception as e:
        logger.error(f"Failed to read file: {e}")
        return None

# ----------------------
# Manual SharePoint Crawler
# ----------------------
def crawl_drive_items(access_token, drive_id, parent_id=None, file_extension=None):
    logger.info(f"Starting manual crawl for drive {drive_id} with filter: {file_extension}")
    headers = {"Authorization": f"Bearer {access_token}"}
    url = f"{GRAPH_URL}/drives/{drive_id}/items/{parent_id}/children" if parent_id else f"{GRAPH_URL}/drives/{drive_id}/root/children"

    results = []
    while url:
        response = requests.get(url, headers=headers)
        if response.status_code != 200:
            logger.error(f"Failed to crawl drive: {response.status_code} - {response.text}")
            break

        data = response.json()
        for item in data.get("value", []):
            if "file" in item:
                if file_extension is None or item["name"].lower().endswith(file_extension.lower()):
                    results.append({
                        "name": item["name"],
                        "id": item["id"],
                        "webUrl": item.get("webUrl"),
                        "driveId": item.get("parentReference", {}).get("driveId")
                    })
            elif "folder" in item:
                # Recursive crawl inside subfolder
                child_id = item["id"]
                child_items = crawl_drive_items(access_token, drive_id, child_id, file_extension)
                results.extend(child_items)

        # Pagination
        url = data.get("@odata.nextLink")

    return results

def resolve_sharepoint_site_id(site_hostname, site_path, access_token):
    """
    Resolves a SharePoint site ID from its hostname and path.
    Example: site_hostname='contoso.sharepoint.com', site_path='/sites/YourSiteName'
    """
    # Ensure site_path starts with a single leading slash
    if not site_path.startswith('/'):
        site_path = '/' + site_path
    url = f"https://graph.microsoft.com/v1.0/sites/{site_hostname}:{site_path}"
    headers = {"Authorization": f"Bearer {access_token}"}
    response = requests.get(url, headers=headers)
    if response.status_code == 200:
        return response.json().get("id")
    else:
        logger.error(f"Could not resolve site: {site_hostname}:{site_path} | {response.status_code} {response.text}")
        return None
